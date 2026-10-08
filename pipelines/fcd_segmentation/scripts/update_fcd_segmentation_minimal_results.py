#!/usr/bin/env python3

from __future__ import annotations

import argparse
import json
import math
import pickle
import shutil
import statistics
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


DEFAULT_WORKBOOK = Path("/nfs/turbo/umms-wilms1/FOMO/inference_outputs/FOMO300K_results_v4.xlsx")
RESULT_ROOT = Path("/nfs/turbo/umms-wilms1/FOMO/inference_outputs")
TASK_DIR = Path(
    "/nfs/turbo/umms-wilms1/FOMO/Data/openneuro_task10_t1flairseg_1mm_20260507/"
    "Task010_OpenNeuro_ds004199_FCDSeg_T1FLAIR_1mm"
)

FLAGGED_QC_CASES = {
    "ds004199seg_sub-00010",
    "ds004199seg_sub-00016",
    "ds004199seg_sub-00050",
    "ds004199seg_sub-00055",
    "ds004199seg_sub-00058",
    "ds004199seg_sub-00063",
    "ds004199seg_sub-00064",
    "ds004199seg_sub-00078",
    "ds004199seg_sub-00122",
}

SHEETS = ["FCD_Seg_Metrics", "FCD_Seg_Predictions", "FCD_Seg_Statistics"]
EXTRA_REMOVE = [
    "FCD_Seg_Master",
    "FCD_Seg_ModelInfo",
    "FCD_Seg_UNetXL_Audit",
    "UNetXL_Settings_Audit",
    "FCD_Seg_PerFold",
    "FCD_Seg_PostprocAudit",
]

MODEL_GROUP = {
    "Ours-23M": "Ours",
    "Ours-23M T1+FLAIR": "Ours",
    "UNet-B": "Baseline",
    "UNet-XL": "Baseline",
    "MMUNetVAE": "Baseline",
}

SPECS = [
    {
        "cohort": "Original full 85",
        "input": "T1+FLAIR",
        "model": "Ours-23M",
        "expected_folds": 5,
        "patterns": [
            "task10_balancedsplit_fixedpp_nsd_reeval_20260522/cleandift_s23_balancedsplit_p192ftv_all_v1_fold*_test.json",
            "task10_balancedsplit_20260512/cleandift_s23_balancedsplit_p192ftv_all_v1_fold*_test.json",
        ],
    },
    {
        "cohort": "Original full 85",
        "input": "T1+FLAIR",
        "model": "UNet-B",
        "expected_folds": 5,
        "patterns": [
            "task10_balancedsplit_fixedpp_nsd_reeval_20260522/unet_b_balancedsplit_unetb_basicdicece_v2_fold*_test.json",
            "task10_balancedsplit_20260512/unet_b_balancedsplit_unetb_basicdicece_v2_fold*_test.json",
        ],
    },
    {
        "cohort": "Original full 85",
        "input": "T1+FLAIR",
        "model": "UNet-XL",
        "expected_folds": 5,
        "patterns": [
            "task10_balancedsplit_fixedpp_nsd_reeval_20260522/unet_xl_balancedsplit_unetxl_basicdicece_v2_fold*_test.json",
            "task10_balancedsplit_20260512/unet_xl_balancedsplit_unetxl_basicdicece_v2_fold*_test.json",
        ],
    },
    {
        "cohort": "Original full 85",
        "input": "T1+FLAIR",
        "model": "MMUNetVAE",
        "expected_folds": 5,
        "patterns": [
            "task10_balancedsplit_fixedpp_nsd_reeval_20260522/mmunetvae_balancedsplit_mmu64b4_noneftv_v1_fold*_test.json",
            "task10_balancedsplit_20260513/mmunetvae_balancedsplit_mmu64b4_noneftv_v1_fold*_test.json",
        ],
    },
    {
        "cohort": "QC76 flagged excluded",
        "input": "FLAIR only",
        "model": "Ours-23M",
        "expected_folds": 5,
        "patterns": [
            "task10_qc76_argmax_no_tta_20260520/"
            "cleandift_s23_qc76_flaironly_b4tb25_imagerefine_focaltversky_best_argmax_notta_fold*_test.json"
        ],
    },
    {
        "cohort": "QC76 flagged excluded",
        "input": "FLAIR only",
        "model": "UNet-B",
        "expected_folds": 5,
        "patterns": [
            "task10_qc76_flaironly_modalitymatched_20260521/"
            "unet_b_qc76_flaironly_unetb_b4_all_dicece_100b_argmax_v1_fold*_test.json"
        ],
    },
    {
        "cohort": "QC76 flagged excluded",
        "input": "FLAIR only",
        "model": "UNet-XL",
        "expected_folds": 5,
        "placeholder_split": "nested_balanced_qc76_kfold",
        "patterns": [
            "task10_qc76_flaironly_modalitymatched_20260521/"
            "unet_xl_qc76_flaironly_unetxl_b4_all_dicece_100b_argmax_v1_fold*_test.json"
        ],
    },
    {
        "cohort": "QC76 flagged excluded",
        "input": "FLAIR only",
        "model": "MMUNetVAE",
        "expected_folds": 5,
        "patterns": [
            "task10_qc76_flaironly_modalitymatched_20260521/"
            "mmunetvae_qc76_flaironly_mmunetvae_b4_none_dicece_100b_argmax_v1_fold*_test.json"
        ],
    },
    {
        "cohort": "QC76 flagged excluded",
        "input": "T1+FLAIR",
        "model": "Ours-23M T1+FLAIR",
        "expected_folds": 5,
        "patterns": [
            "task10_qc76_ours_t1flair_20260521/"
            "cleandift_s23_qc76_t1flair_ours_b1acc4_tb25_earlyfusion_ftv_argmax_v1_fold*_test.json"
        ],
    },
]


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser()
    p.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    return p.parse_args()


def result_files(patterns: list[str], expected_folds: int | None = None) -> list[Path]:
    pattern_results: list[dict[int, Path]] = []
    for pattern in patterns:
        selected: dict[int, Path] = {}
        for path in sorted(RESULT_ROOT.glob(pattern)):
            payload = json.loads(path.read_text())
            fold = int(payload.get("fold_idx", path.stem.split("_fold")[-1].split("_")[0]))
            selected[fold] = path
        pattern_results.append(selected)

    if len(patterns) > 1 and expected_folds is not None:
        for selected in pattern_results:
            if len(selected) >= expected_folds:
                return [selected[k] for k in sorted(selected)]

    selected: dict[int, Path] = {}
    for per_pattern in pattern_results:
        for fold, path in sorted(per_pattern.items()):
            if fold not in selected:
                selected[fold] = path
    return [selected[k] for k in sorted(selected)]


def load_predictions(spec: dict[str, Any]) -> tuple[list[dict[str, Any]], list[int]]:
    rows: list[dict[str, Any]] = []
    folds: list[int] = []
    for path in result_files(spec["patterns"], int(spec["expected_folds"])):
        payload = json.loads(path.read_text())
        fold = int(payload["fold_idx"])
        folds.append(fold)
        preds = payload.get("predictions", payload.get("subjects", []))
        status = "Complete" if "test_mean_dice" in payload or "mean_dice" in payload else "Unknown"
        for pred in preds:
            rows.append(
                {
                    "cohort": spec["cohort"],
                    "group": MODEL_GROUP[spec["model"]],
                    "model": spec["model"],
                    "input": spec["input"],
                    "status": status,
                    "fold": fold,
                    "subject": pred["id"],
                    "dice": float(pred["dice"]),
                    "nsd": _optional_float(pred.get("nsd")),
                    "gt_voxels": pred.get("gt_voxels"),
                    "pred_voxels": pred.get("pred_voxels"),
                    "flagged_qc": "YES" if pred["id"] in FLAGGED_QC_CASES else "NO",
                    "result_json": str(path),
                }
            )
    return rows, folds


def _optional_float(value: Any) -> float | None:
    if value is None:
        return None
    value = float(value)
    if math.isnan(value):
        return None
    return value


def pending_rows(spec: dict[str, Any], folds: list[int]) -> list[dict[str, Any]]:
    split_name = spec.get("placeholder_split")
    if not split_name:
        return []
    missing = sorted(set(range(int(spec["expected_folds"]))) - set(folds))
    if not missing:
        return []
    splits = pickle.load(open(TASK_DIR / "splits.pkl", "rb"))
    fold_defs = splits[split_name][int(spec["expected_folds"])]
    rows = []
    for fold in missing:
        for subject in sorted(fold_defs[fold]["test"]):
            rows.append(
                {
                    "cohort": spec["cohort"],
                    "group": MODEL_GROUP[spec["model"]],
                    "model": spec["model"],
                    "input": spec["input"],
                    "status": "Pending",
                    "fold": fold,
                    "subject": subject,
                    "dice": None,
                    "nsd": None,
                    "gt_voxels": None,
                    "pred_voxels": None,
                    "flagged_qc": "YES" if subject in FLAGGED_QC_CASES else "NO",
                    "result_json": "",
                }
            )
    return rows


def collect_rows() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    predictions = []
    metrics = []
    for spec in SPECS:
        pred_rows, folds = load_predictions(spec)
        pred_rows.extend(pending_rows(spec, folds))
        predictions.extend(pred_rows)
        complete = [row for row in pred_rows if row["status"] == "Complete" and row["dice"] is not None]
        dice = [row["dice"] for row in complete]
        nsd = [row["nsd"] for row in complete if row["nsd"] is not None]
        subjects = sorted({row["subject"] for row in complete})
        status = "Complete" if len(set(folds)) >= int(spec["expected_folds"]) else "Pending"
        metrics.append(
            {
                "cohort": spec["cohort"],
                "group": MODEL_GROUP[spec["model"]],
                "model": spec["model"],
                "input": spec["input"],
                "status": status,
                "n_subjects": len(subjects),
                "n_folds": len(set(folds)),
                "expected_folds": int(spec["expected_folds"]),
                "mean_dice": statistics.mean(dice) if dice else None,
                "std_dice": statistics.pstdev(dice) if len(dice) > 1 else 0.0 if dice else None,
                "median_dice": statistics.median(dice) if dice else None,
                "mean_nsd": statistics.mean(nsd) if nsd else None,
                "std_nsd": statistics.pstdev(nsd) if len(nsd) > 1 else 0.0 if nsd else None,
                "median_nsd": statistics.median(nsd) if nsd else None,
            }
        )
    predictions.sort(key=lambda r: (r["cohort"], r["input"], r["group"], r["model"], r["fold"], r["subject"]))
    metrics.sort(key=lambda r: (r["cohort"], r["input"], r["group"], r["model"]))
    return metrics, predictions


def paired_stats(predictions: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    by_group: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for row in predictions:
        if row["status"] == "Complete":
            by_group.setdefault((row["cohort"], row["input"]), []).append(row)
    for (cohort, input_name), items in sorted(by_group.items()):
        models = sorted({row["model"] for row in items})
        ours = "Ours-23M" if "Ours-23M" in models else None
        if ours is None:
            continue
        for other in models:
            if other == ours or other.startswith("Ours-23M T1+FLAIR"):
                continue
            for metric in ("dice", "nsd"):
                left = {row["subject"]: row[metric] for row in items if row["model"] == ours and row[metric] is not None}
                right = {row["subject"]: row[metric] for row in items if row["model"] == other and row[metric] is not None}
                subjects = sorted(set(left) & set(right))
                if not subjects:
                    continue
                a = [float(left[s]) for s in subjects]
                b = [float(right[s]) for s in subjects]
                diffs = [x - y for x, y in zip(a, b)]
                rows.append(
                    {
                        "cohort": cohort,
                        "input": input_name,
                        "metric": metric.upper(),
                        "comparison": f"{ours} - {other}",
                        "n": len(subjects),
                        "mean_ours": statistics.mean(a),
                        "mean_other": statistics.mean(b),
                        "mean_difference": statistics.mean(diffs),
                        "wilcoxon_p": wilcoxon_p(a, b),
                    }
                )
    return rows


def wilcoxon_p(a: list[float], b: list[float]) -> float | None:
    try:
        from scipy.stats import wilcoxon
    except Exception:
        return None
    try:
        stat = wilcoxon(np.asarray(a), np.asarray(b), zero_method="zsplit")
    except ValueError:
        return None
    return float(stat.pvalue)


def append_title(ws, title: str, subtitle: str, width: int) -> None:
    ws.append([title])
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.append([subtitle])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws.append([])


def append_header(ws, header: list[str]) -> None:
    ws.append(header)
    fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")


def autosize(ws) -> None:
    for column_cells in ws.columns:
        first_cell = next((cell for cell in column_cells if hasattr(cell, "column_letter")), None)
        if first_cell is None:
            continue
        letter = first_cell.column_letter
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 80))
        ws.column_dimensions[letter].width = max(10, min(max_len + 2, 60))


def write_workbook(workbook: Path, metrics: list[dict[str, Any]], predictions: list[dict[str, Any]]) -> None:
    backup = workbook.with_suffix(f".backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    shutil.copy2(workbook, backup)
    wb = load_workbook(workbook)
    for sheet in SHEETS + EXTRA_REMOVE:
        if sheet in wb.sheetnames:
            del wb[sheet]

    ws = wb.create_sheet("FCD_Seg_Metrics")
    append_title(ws, "FCD segmentation metrics", "Held-out test performance. NSD is listed next to Dice.", 12)
    append_header(
        ws,
        [
            "Cohort",
            "Group",
            "Model",
            "Input",
            "Status",
            "N Subjects",
            "N Folds",
            "Expected Folds",
            "Mean Dice",
            "Std Dice",
            "Median Dice",
            "Mean NSD",
            "Std NSD",
            "Median NSD",
        ],
    )
    for row in metrics:
        ws.append(
            [
                row["cohort"],
                row["group"],
                row["model"],
                row["input"],
                row["status"],
                row["n_subjects"],
                row["n_folds"],
                row["expected_folds"],
                row["mean_dice"],
                row["std_dice"],
                row["median_dice"],
                row["mean_nsd"],
                row["std_nsd"],
                row["median_nsd"],
            ]
        )
    autosize(ws)

    ws = wb.create_sheet("FCD_Seg_Predictions")
    append_title(ws, "FCD segmentation predictions", "Held-out test case Dice/NSD.", 12)
    append_header(
        ws,
        [
            "Cohort",
            "Group",
            "Model",
            "Input",
            "Status",
            "Fold",
            "Subject",
            "Dice",
            "NSD",
            "GT Voxels",
            "Pred Voxels",
            "Flagged QC",
        ],
    )
    for row in predictions:
        ws.append(
            [
                row["cohort"],
                row["group"],
                row["model"],
                row["input"],
                row["status"],
                row["fold"],
                row["subject"],
                row["dice"],
                row["nsd"],
                row["gt_voxels"],
                row["pred_voxels"],
                row["flagged_qc"],
            ]
        )
    autosize(ws)

    stats = paired_stats(predictions)
    ws = wb.create_sheet("FCD_Seg_Statistics")
    append_title(ws, "FCD segmentation statistics", "Paired per-case comparisons on shared held-out test subjects.", 9)
    append_header(
        ws,
        [
            "Cohort",
            "Input",
            "Metric",
            "Comparison",
            "N",
            "Mean Ours",
            "Mean Other",
            "Mean Difference",
            "Wilcoxon p",
        ],
    )
    for row in stats:
        ws.append(
            [
                row["cohort"],
                row["input"],
                row["metric"],
                row["comparison"],
                row["n"],
                row["mean_ours"],
                row["mean_other"],
                row["mean_difference"],
                row["wilcoxon_p"],
            ]
        )
    autosize(ws)

    wb.save(workbook)
    print(f"Updated {workbook}")
    print(f"Backup {backup}")


def main() -> None:
    args = parse_args()
    metrics, predictions = collect_rows()
    write_workbook(args.workbook, metrics, predictions)
    for row in metrics:
        print(
            row["cohort"],
            row["input"],
            row["model"],
            row["status"],
            row["n_folds"],
            row["mean_dice"],
            row["mean_nsd"],
        )


if __name__ == "__main__":
    main()
