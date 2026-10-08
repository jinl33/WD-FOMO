#!/usr/bin/env python3

from __future__ import annotations

import json
import math
import re
import statistics
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


WB_PATH = Path(
    "/nfs/turbo/umms-wilms1/FOMO/inference_outputs/"
    "FOMO300K_results_v3_siteagebalanced_20260504.xlsx"
)
RESULT_DIR = Path(
    "/nfs/turbo/umms-wilms1/FOMO/inference_outputs/"
    "task10_balancedsplit_20260512"
)

THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
TITLE_FONT = Font(bold=True, size=13)
HDR_FONT = Font(bold=True, size=11)
FILL_BLUE = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

MODEL_LABELS = {
    "cleandift_s23": "Ours-23M",
    "mmunetvae": "MMUNetVAE",
    "unet_b": "UNet-B",
    "unet_xl": "UNet-XL",
}
GROUP_LABELS = {
    "cleandift_s23": "Ours",
    "mmunetvae": "Baseline",
    "unet_b": "Baseline",
    "unet_xl": "Baseline",
}


def load_json(path: Path) -> dict:
    with open(path) as f:
        return json.load(f)


def section_title(ws, text: str, cols: int) -> None:
    ws.append([text])
    ws.merge_cells(
        start_row=ws.max_row,
        start_column=1,
        end_row=ws.max_row,
        end_column=cols,
    )
    cell = ws.cell(ws.max_row, 1)
    cell.font = TITLE_FONT
    cell.alignment = Alignment(horizontal="left")


def header_row(ws, values: list[str]) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font = HDR_FONT
        cell.fill = FILL_BLUE
        cell.border = THIN
        cell.alignment = Alignment(horizontal="center")


def format_last_row(ws) -> None:
    for cell in ws[ws.max_row]:
        cell.border = THIN
        if isinstance(cell.value, float):
            cell.number_format = "0.0000"


def autosize(ws) -> None:
    widths: dict[str, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            widths[cell.column_letter] = max(
                widths.get(cell.column_letter, 0), len(str(cell.value))
            )
    for col, width in widths.items():
        ws.column_dimensions[col].width = min(max(width + 2, 12), 42)


def describe_setting(path: Path) -> str:
    stem = path.stem
    if "balancedsplit_p192ftv_all" in stem:
        return "balanced nested CV, full 192x256x192, aug=all, focaltversky"
    if "balancedsplit_p128x160x128" in stem:
        return "balanced fold-0 screen, 128x160x128, aug=all, focaltversky"
    if "balancedsplit_p160x192x160" in stem:
        return "balanced fold-0 screen, 160x192x160, aug=all, focaltversky"
    if "balancedsplit_mmu_noneftv" in stem:
        return "balanced nested CV, 96x96x96, aug=none, focaltversky"
    if "balancedsplit_unetb_basicdicece" in stem:
        return "balanced nested CV, 96x96x96, aug=basic, dicece"
    if "balancedsplit_unetxl_basicdicece" in stem:
        return "balanced nested CV, 96x96x96, aug=basic, dicece"
    return stem


def stable_setting_key(path: Path) -> str:
    return re.sub(r"_fold\d+$", "", path.stem.replace("_test", ""))


def corr(xs: list[float], ys: list[float]) -> float | None:
    if len(xs) < 2 or len(xs) != len(ys):
        return None
    mx = sum(xs) / len(xs)
    my = sum(ys) / len(ys)
    num = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    den = math.sqrt(sum((x - mx) ** 2 for x in xs) * sum((y - my) ** 2 for y in ys))
    return num / den if den else None


def collect_rows() -> tuple[list[dict], list[dict], list[dict]]:
    rows = []
    subject_rows = []
    for path in sorted(RESULT_DIR.glob("*_fold*_test.json")):
        payload = load_json(path)
        model_name = payload["model_name"]
        best_pp = payload.get("best_postprocess") or {}
        test_mean = payload.get("test_mean_dice", payload.get("mean_dice"))
        test_std = payload.get("test_std_dice", payload.get("std_dice"))
        test_median = payload.get("test_median_dice", payload.get("median_dice"))
        row = {
            "setting_key": stable_setting_key(path),
            "group": GROUP_LABELS.get(model_name, ""),
            "model": MODEL_LABELS.get(model_name, model_name),
            "model_name": model_name,
            "setting": describe_setting(path),
            "fold_idx": payload.get("fold_idx", payload.get("fold")),
            "n_val_subjects": payload.get("n_val_subjects"),
            "n_test_subjects": payload.get("n_test_subjects", payload.get("n_subjects")),
            "test_mean_dice": test_mean,
            "test_std_dice": test_std,
            "test_median_dice": test_median,
            "threshold": best_pp.get("threshold"),
            "min_component_voxels": best_pp.get("min_component_voxels"),
            "keep_largest": best_pp.get("keep_largest"),
            "val_mean_dice": best_pp.get("val_mean_dice"),
            "val_std_dice": best_pp.get("val_std_dice"),
            "val_median_dice": best_pp.get("val_median_dice"),
            "path": str(path),
        }
        rows.append(row)
        for pred in payload.get("predictions", []):
            subject_rows.append(
                {
                    "group": row["group"],
                    "model": row["model"],
                    "setting": row["setting"],
                    "fold_idx": row["fold_idx"],
                    "id": pred.get("id"),
                    "dice": pred.get("dice"),
                    "gt_voxels": pred.get("gt_voxels"),
                    "pred_voxels": pred.get("pred_voxels"),
                }
            )

    rows.sort(key=lambda r: (r["model"], r["setting"], r["fold_idx"]))
    subject_rows.sort(key=lambda r: (r["model"], r["setting"], r["fold_idx"], r["id"]))

    grouped: dict[str, list[dict]] = {}
    for row in rows:
        grouped.setdefault(row["setting_key"], []).append(row)

    cv_rows = []
    for items in grouped.values():
        items = sorted(items, key=lambda r: r["fold_idx"])
        fold_means = [float(item["test_mean_dice"]) for item in items]
        val_means = [
            float(item["val_mean_dice"])
            for item in items
            if item["val_mean_dice"] is not None
        ]
        val_test_corr = corr(val_means, fold_means) if len(val_means) == len(fold_means) else None
        cv_rows.append(
            {
                "group": items[0]["group"],
                "model": items[0]["model"],
                "setting": items[0]["setting"],
                "n_folds": len(items),
                "folds": ",".join(str(item["fold_idx"]) for item in items),
                "mean_fold_dice": statistics.mean(fold_means),
                "std_fold_dice": statistics.stdev(fold_means) if len(fold_means) > 1 else 0.0,
                "median_fold_dice": statistics.median(fold_means),
                "best_fold_dice": max(fold_means),
                "worst_fold_dice": min(fold_means),
                "val_test_corr": val_test_corr,
            }
        )
    cv_rows.sort(key=lambda r: (-r["mean_fold_dice"], r["model"], r["setting"]))
    return cv_rows, rows, subject_rows


def main() -> None:
    if not WB_PATH.exists():
        raise SystemExit(f"Missing workbook: {WB_PATH}")
    cv_rows, rows, subject_rows = collect_rows()
    if not rows:
        raise SystemExit(f"No balanced result JSONs found in {RESULT_DIR}")

    wb = load_workbook(WB_PATH)
    for name in [
        "FCD_Seg_Balanced_CV",
        "FCD_Seg_Balanced_Folds",
        "FCD_Seg_Balanced_Subjects",
    ]:
        if name in wb.sheetnames:
            del wb[name]

    ws = wb.create_sheet("FCD_Seg_Balanced_CV")
    section_title(ws, "ds004199 FCD Segmentation - Balanced Nested CV Summary", 11)
    ws.append([])
    header_row(
        ws,
        [
            "Group",
            "Model",
            "Setting",
            "N Folds",
            "Folds",
            "Mean Fold Dice",
            "Std Fold Dice",
            "Median Fold Dice",
            "Best Fold Dice",
            "Worst Fold Dice",
            "Val/Test Corr",
        ],
    )
    for row in cv_rows:
        ws.append(
            [
                row["group"],
                row["model"],
                row["setting"],
                row["n_folds"],
                row["folds"],
                row["mean_fold_dice"],
                row["std_fold_dice"],
                row["median_fold_dice"],
                row["best_fold_dice"],
                row["worst_fold_dice"],
                row["val_test_corr"],
            ]
        )
        format_last_row(ws)
    ws.freeze_panes = "A4"
    autosize(ws)

    ws = wb.create_sheet("FCD_Seg_Balanced_Folds")
    section_title(ws, "ds004199 FCD Segmentation - Balanced Per-Fold Summary", 16)
    ws.append([])
    header_row(
        ws,
        [
            "Group",
            "Model",
            "Setting",
            "Fold",
            "N Val",
            "N Test",
            "Test Mean Dice",
            "Test Std Dice",
            "Test Median Dice",
            "Threshold",
            "Min Component Voxels",
            "Keep Largest",
            "Val Mean Dice",
            "Val Std Dice",
            "Val Median Dice",
            "Result file",
        ],
    )
    for row in rows:
        ws.append(
            [
                row["group"],
                row["model"],
                row["setting"],
                row["fold_idx"],
                row["n_val_subjects"],
                row["n_test_subjects"],
                row["test_mean_dice"],
                row["test_std_dice"],
                row["test_median_dice"],
                row["threshold"],
                row["min_component_voxels"],
                int(bool(row["keep_largest"])) if row["keep_largest"] is not None else None,
                row["val_mean_dice"],
                row["val_std_dice"],
                row["val_median_dice"],
                row["path"],
            ]
        )
        format_last_row(ws)
    ws.freeze_panes = "A4"
    autosize(ws)

    ws = wb.create_sheet("FCD_Seg_Balanced_Subjects")
    section_title(ws, "ds004199 FCD Segmentation - Balanced Per-Subject Dice", 8)
    ws.append([])
    header_row(
        ws,
        ["Group", "Model", "Setting", "Fold", "Subject ID", "Dice", "GT Voxels", "Pred Voxels"],
    )
    for row in subject_rows:
        ws.append(
            [
                row["group"],
                row["model"],
                row["setting"],
                row["fold_idx"],
                row["id"],
                row["dice"],
                row["gt_voxels"],
                row["pred_voxels"],
            ]
        )
        format_last_row(ws)
    ws.freeze_panes = "A4"
    autosize(ws)

    wb.save(WB_PATH)
    print(f"Updated workbook: {WB_PATH}")
    for row in cv_rows:
        print(
            f"{row['model']}: n={row['n_folds']} "
            f"mean={row['mean_fold_dice']:.4f} sd={row['std_fold_dice']:.4f} "
            f"{row['setting']}"
        )


if __name__ == "__main__":
    main()
