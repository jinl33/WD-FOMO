#!/usr/bin/env python3

from __future__ import annotations

import argparse
import json
import math
import re
import statistics
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


DEFAULT_WORKBOOK = Path(
    "/nfs/turbo/umms-wilms1/FOMO/inference_outputs/"
    "FOMO300K_results_v3_siteagebalanced_20260504.xlsx"
)
DEFAULT_RESULT_ROOT = Path("/nfs/turbo/umms-wilms1/FOMO/inference_outputs")
MASTER_SHEET = "FCD_Seg_Master"
REPLACED_SHEETS = [
    "FCD_Seg_Metrics",
    "FCD_Seg_Predictions",
    "FCD_Seg_Statistics",
    "FCD_Seg_ModelInfo",
    "FCD_Seg_UNetXL_Audit",
    "FCD_Seg_PerFold",
    "FCD_Seg_PostprocAudit",
]
FLAGGED_QC_CASES = {
    "ds004199seg_sub-00010",
    "ds004199seg_sub-00016",
    "ds004199seg_sub-00050",
    "ds004199seg_sub-00055",
    "ds004199seg_sub-00058",
    "ds004199seg_sub-00063",
    "ds004199seg_sub-00064",
    "ds004199seg_sub-00078",
    "ds004199seg_sub-00122",
}

THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")

MODEL_DISPLAY = {
    "cleandift_s23": "Ours-23M",
    "unet_b": "UNet-B",
    "unet_xl": "UNet-XL",
    "mmunetvae": "MMUNetVAE",
}

MODEL_INFO = {
    "cleandift_s23": {
        "group": "Ours",
        "architecture": (
            "Pretrained time-independent CleanDIFT wavelet diffusion backbone. "
            "Segmentation uses backbone multi-scale features with a trainable 3D decoder."
        ),
        "input": (
            "Current QC76 result uses FLAIR only. The volume is converted with a two-level "
            "3D DWPT into 64 wavelet subbands before the pretrained backbone."
        ),
        "resolution": "Full volume 192 x 256 x 192, batch size 4.",
        "decoder": "Image-refinement segmentation head with skip features from the CleanDIFT backbone.",
    },
    "unet_b": {
        "group": "Baseline",
        "architecture": (
            "FOMO25 baseline spatial-domain 3D U-Net encoder-decoder using direct image channels, "
            "standard skip connections, and 32 starting filters."
        ),
        "input": "T1w + FLAIR image channels.",
        "resolution": "Patch-based training and inference with 96 x 96 x 96 patches. Main QC76 rows use batch size 4; baseline-default sensitivity rows use batch size 2.",
        "decoder": "Native U-Net decoder from the baseline codebase.",
    },
    "unet_xl": {
        "group": "Baseline",
        "architecture": (
            "Larger FOMO25 baseline spatial-domain 3D U-Net encoder-decoder using direct image "
            "channels, standard skip connections, and 64 starting filters."
        ),
        "input": "T1w + FLAIR image channels.",
        "resolution": "Patch-based training and inference with 96 x 96 x 96 patches. Main QC76 rows use batch size 4; baseline-default sensitivity rows use batch size 2.",
        "decoder": "Native U-Net decoder from the baseline codebase. The fine-tuned model has about 4x the trainable parameters of UNet-B.",
    },
    "mmunetvae": {
        "group": "Baseline",
        "architecture": (
            "Multimodal U-Net/VAE-style baseline with modality-specific feature extraction, "
            "feature fusion, and a U-Net decoder."
        ),
        "input": "T1w + FLAIR image channels.",
        "resolution": "Patch-based training and inference with 64 x 64 x 64 patches, batch size 4.",
        "decoder": "Native MMUNetVAE segmentation decoder.",
    },
}

COMMON_SEGMENTATION_SETUP = (
    "FCD segmentation uses ds004199 T1w+FLAIR. T1w is the reference image; FLAIR and ROI are "
    "affine/header-resampled into the T1w grid, with linear interpolation for FLAIR and "
    "nearest-neighbor interpolation for ROI. T1w is skull-stripped with SynthStrip, and the "
    "T1-derived brain mask is applied to aligned FLAIR before shared preprocessing."
)

COMPARISON_SPECS = [
    {
        "cohort": "Original full 85",
        "model": "cleandift_s23",
        "expected_folds": 5,
        "keys": ["task10_balancedsplit_20260512/cleandift_s23_balancedsplit_p192ftv_all_v1"],
        "note": "Original full-85 5-fold result.",
    },
    {
        "cohort": "Original full 85",
        "model": "unet_b",
        "expected_folds": 5,
        "keys": ["task10_balancedsplit_20260512/unet_b_balancedsplit_unetb_basicdicece_v2"],
        "note": "Original full-85 5-fold result.",
    },
    {
        "cohort": "Original full 85",
        "model": "unet_xl",
        "expected_folds": 5,
        "keys": ["task10_balancedsplit_20260512/unet_xl_balancedsplit_unetxl_basicdicece_v2"],
        "note": "Original full-85 5-fold result.",
    },
    {
        "cohort": "Original full 85",
        "model": "mmunetvae",
        "expected_folds": 5,
        "keys": ["task10_balancedsplit_20260513/mmunetvae_balancedsplit_mmu64b4_noneftv_v1"],
        "note": "Original full-85 5-fold result.",
    },
    {
        "cohort": "QC76 flagged excluded",
        "model": "cleandift_s23",
        "expected_folds": 5,
        "eval_mode": "tuned",
        "keys": [
            "task10_qc76_20260517/"
            "cleandift_s23_qc76_flaironly_b4tb25_imagerefine_focaltversky_fullLR_best_nw2_v2",
        ],
        "note": "QC76 final run using best.ckpt and validation-tuned postprocessing.",
    },
    {
        "cohort": "QC76 flagged excluded",
        "model": "unet_b",
        "expected_folds": 5,
        "eval_mode": "argmax_no_tune",
        "keys": [
            "task10_qc76_unetb_reseed_20260519/"
            "unet_b_qc76_unetb_b4_all_dicece_best_100b_seed20260519_v1",
        ],
        "note": (
            "QC76 final baseline run using best.ckpt, aug=all, 100 train batches/epoch, "
            "and baseline-repo-style argmax inference without tuned postprocessing."
        ),
    },
    {
        "cohort": "QC76 flagged excluded",
        "model": "unet_xl",
        "expected_folds": 5,
        "eval_mode": "argmax_no_tune",
        "keys": [
            "task10_qc76_unetxl_reseed_20260519/"
            "unet_xl_qc76_unetxl_b4_all_dicece_best_100b_seed20260519_v1",
        ],
        "note": (
            "QC76 final baseline run using best.ckpt, aug=all, 100 train batches/epoch, "
            "and baseline-repo-style argmax inference without tuned postprocessing."
        ),
    },
    {
        "cohort": "QC76 flagged excluded",
        "model": "mmunetvae",
        "expected_folds": 5,
        "eval_mode": "tuned",
        "keys": [
            "task10_qc76_baselines_20260517/mmunetvae_qc76_mmunetvae_b4_none_dicece_best_nw2_v2_wilms"
        ],
        "note": "QC76 final result using best.ckpt and tuned postprocessing.",
    },
    {
        "cohort": "QC76 baseline-default sensitivity",
        "model": "unet_b",
        "expected_folds": 5,
        "eval_mode": "argmax_no_tune",
        "keys": [
            "task10_qc76_baseline_default_20260520/"
            "unet_b_qc76_unet_b_b2_basic_dicece_best_500b_seed20260520_baselinedef_v2",
        ],
        "note": (
            "Sensitivity run matching the baseline-codebase fine-tuning example more closely: "
            "batch size 2, aug=basic, DiceCE, 500 max epochs with early stopping, 100 train "
            "batches/epoch, best.ckpt, argmax inference."
        ),
    },
    {
        "cohort": "QC76 baseline-default sensitivity",
        "model": "unet_xl",
        "expected_folds": 5,
        "eval_mode": "argmax_no_tune",
        "keys": [
            "task10_qc76_baseline_default_20260520/"
            "unet_xl_qc76_unet_xl_b2_basic_dicece_best_500b_seed20260520_baselinedef_v2",
        ],
        "note": (
            "Sensitivity run matching the baseline-codebase fine-tuning example more closely: "
            "batch size 2, aug=basic, DiceCE, 500 max epochs with early stopping, 100 train "
            "batches/epoch, best.ckpt, argmax inference."
        ),
    },
]


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser()
    p.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    p.add_argument("--result_root", type=Path, default=DEFAULT_RESULT_ROOT)
    p.add_argument(
        "--result_glob",
        action="append",
        default=None,
        help="Result directory glob under result_root. Can be passed multiple times.",
    )
    return p.parse_args()


def result_globs(args: argparse.Namespace) -> list[str]:
    return args.result_glob or [
        "task10_*",
        "task2_meningioma_*",
    ]


def load_payload(path: Path) -> dict[str, Any] | None:
    if path.name.endswith("_timing.json"):
        return None
    try:
        payload = json.loads(path.read_text())
    except Exception:
        return None
    if "model_name" not in payload or "predictions" not in payload:
        return None
    if payload.get("test_mean_dice") is None:
        return None
    return payload


def load_timing(path: Path) -> dict[str, Any]:
    timing_path = path.with_name(path.name.replace("_test.json", "_timing.json"))
    if not timing_path.exists():
        return {}
    try:
        return json.loads(timing_path.read_text())
    except Exception:
        return {}


def stable_key(path: Path) -> str:
    stem = path.stem
    stem = stem.replace("_recovered_best_tuned_test", "_test")
    stem = stem.replace("_recovered_best_argmax_test", "_argmax_test")
    stem = stem.replace("_tuned_fold", "_fold")
    stem = re.sub(r"_fold\d+_last_test$", "_last_test", stem)
    stem = re.sub(r"_fold\d+", "", stem)
    stem = re.sub(r"_test$", "", stem)
    try:
        parent = path.parent.relative_to(DEFAULT_RESULT_ROOT)
    except ValueError:
        parent = Path(path.parent.name)
    return f"{parent}/{stem}"


def model_group(model_name: str) -> str:
    if model_name.startswith("cleandift"):
        return "Ours"
    if model_name in {"unet_b", "unet_xl"}:
        return "UNet baseline"
    if model_name == "mmunetvae":
        return "MMUNetVAE baseline"
    return "Other"


def checkpoint_label(payload: dict[str, Any]) -> str:
    ckpt = Path(str(payload.get("checkpoint", ""))).name
    if ckpt == "best.ckpt" or "best" in ckpt.lower():
        return "best"
    if ckpt == "last.ckpt" or "last" in ckpt.lower():
        return "last"
    return ckpt or "unknown"


def short_setting(key: str, payload: dict[str, Any]) -> str:
    parts = [payload.get("model_name", "model")]
    seg_head = payload.get("seg_head_variant")
    if (
        not seg_head
        and payload.get("model_name") in {"unet_b", "unet_xl", "mmunetvae"}
        and payload.get("evaluation_mode") == "tuned"
        and (
            "task10_qc76_baselines_20260517" in key
            or "task10_qc76_baselines_fulltrain_20260518" in key
        )
    ):
        seg_head = "earlyfusion"
    if seg_head:
        parts.append(f"head={seg_head}")
    if payload.get("input_modality_indices") is not None:
        parts.append(f"modalities={payload['input_modality_indices']}")
    if payload.get("evaluation_mode"):
        parts.append(str(payload["evaluation_mode"]))
    parts.append(key.split("/", 1)[0])
    stem = key.split("/", 1)[1] if "/" in key else key
    if any(token in stem for token in ("snapshot", "currentbest", "evalonly")):
        parts.append(stem)
    return "; ".join(str(x) for x in parts)


def cohort_label(row_or_payload: dict[str, Any], key: str = "") -> str:
    split_method = str(row_or_payload.get("split_method") or "")
    path_key = key.lower()
    if "task2" in path_key or "meningioma" in path_key:
        return "Task2 meningioma"
    if "qc76" in split_method.lower() or "qc76" in path_key:
        return "FCD QC76, flagged excluded"
    return "FCD full 85, flagged included"


def cohort_priority(cohort: str) -> int:
    if cohort.startswith("FCD QC76"):
        return 0
    if cohort.startswith("FCD full"):
        return 1
    return 2


def mean_or_none(values: list[float]) -> float | None:
    return statistics.mean(values) if values else None


def percentile_or_none(values: list[float], percentile: float) -> float | None:
    if not values:
        return None
    xs = sorted(float(x) for x in values)
    pos = (len(xs) - 1) * percentile / 100.0
    lo = int(pos)
    hi = min(lo + 1, len(xs) - 1)
    if lo == hi:
        return xs[lo]
    frac = pos - lo
    return xs[lo] * (1.0 - frac) + xs[hi] * frac


def fold_row(path: Path, payload: dict[str, Any]) -> dict[str, Any]:
    timing = load_timing(path)
    key = stable_key(path)
    checkpoint = checkpoint_label(payload)
    model_name = str(payload.get("model_name"))
    best_monitor = timing.get("best_checkpoint_monitor")
    best_mode = timing.get("best_checkpoint_mode")
    if checkpoint == "best" and not best_monitor:
        if model_name.startswith("cleandift"):
            best_monitor = "val/soft_dice"
        elif model_name in {"unet_b", "unet_xl", "mmunetvae"}:
            best_monitor = "val/dice"
        best_mode = best_mode or "max"
    if (
        payload.get("model_name") in {"unet_b", "unet_xl", "mmunetvae"}
        and payload.get("evaluation_mode") == "tuned"
        and (
            "task10_qc76_baselines_20260517" in key
            or "task10_qc76_baselines_fulltrain_20260518" in key
        )
    ):
        best_monitor = best_monitor or "val/dice"
        best_mode = best_mode or "max"
    preds = payload.get("predictions", [])
    flagged = [p for p in preds if p.get("id") in FLAGGED_QC_CASES]
    unflagged = [p for p in preds if p.get("id") not in FLAGGED_QC_CASES]
    case_dices = [float(p.get("dice", 0.0)) for p in preds]
    flagged_dice = [float(p.get("dice", 0.0)) for p in flagged]
    unflagged_dice = [float(p.get("dice", 0.0)) for p in unflagged]
    return {
        "key": key,
        "cohort": cohort_label(payload, key),
        "group": model_group(str(payload.get("model_name"))),
        "model": payload.get("model_name"),
        "setting": short_setting(key, payload),
        "split_method": payload.get("split_method"),
        "split_param": payload.get("split_param"),
        "fold": payload.get("fold_idx"),
        "checkpoint": checkpoint,
        "best_monitor": best_monitor,
        "best_mode": best_mode,
        "eval_mode": payload.get("evaluation_mode"),
        "test_dice": float(payload.get("test_mean_dice")),
        "test_nsd": payload.get("test_mean_nsd"),
        "case_dices": case_dices,
        "case_median_dice": statistics.median(case_dices) if case_dices else None,
        "case_iqr_low": percentile_or_none(case_dices, 25),
        "case_iqr_high": percentile_or_none(case_dices, 75),
        "zero_dice_cases": sum(1 for x in case_dices if x <= 0.0),
        "dice_ge_0_1_cases": sum(1 for x in case_dices if x >= 0.1),
        "dice_ge_0_2_cases": sum(1 for x in case_dices if x >= 0.2),
        "unflagged_dice": mean_or_none(unflagged_dice),
        "flagged_dice": mean_or_none(flagged_dice),
        "n_flagged": len(flagged),
        "n_test": len(preds),
        "path": str(path),
        "predictions": [
            {
                "id": p.get("id"),
                "dice": float(p.get("dice", 0.0)),
                "nsd": p.get("nsd"),
                "gt_voxels": p.get("gt_voxels"),
                "pred_voxels": p.get("pred_voxels"),
                "flagged_qc": p.get("id") in FLAGGED_QC_CASES,
            }
            for p in preds
        ],
    }


def collect_rows(args: argparse.Namespace) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    seen: set[Path] = set()
    for pattern in result_globs(args):
        for result_dir in sorted(args.result_root.glob(pattern)):
            if not result_dir.is_dir():
                continue
            for path in sorted(result_dir.rglob("*_test.json")):
                if path in seen:
                    continue
                seen.add(path)
                payload = load_payload(path)
                if payload is None:
                    continue
                rows.append(fold_row(path, payload))
    return rows


def summarize(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    for row in rows:
        group_key = (
            row["key"],
            row["cohort"],
            row["group"],
            row["model"],
            row["setting"],
            row["split_method"],
            row["split_param"],
            row["checkpoint"],
            row["best_monitor"],
            row["best_mode"],
            row["eval_mode"],
        )
        grouped.setdefault(group_key, []).append(row)

    summary_rows = []
    for (
        key,
        cohort,
        group,
        model,
        setting,
        split_method,
        split_param,
        ckpt,
        best_monitor,
        best_mode,
        eval_mode,
    ), items in grouped.items():
        items = sorted(items, key=lambda x: (-1 if x["fold"] is None else int(x["fold"])))
        dices = [float(x["test_dice"]) for x in items]
        nsds = [float(x["test_nsd"]) for x in items if x["test_nsd"] is not None]
        case_dices = [
            float(dice)
            for item in items
            for dice in item.get("case_dices", [])
        ]
        unflagged = [float(x["unflagged_dice"]) for x in items if x["unflagged_dice"] is not None]
        flagged = [float(x["flagged_dice"]) for x in items if x["flagged_dice"] is not None]
        summary_rows.append(
            {
                "group": group,
                "cohort": cohort,
                "model": model,
                "setting": setting,
                "split_method": split_method,
                "split_param": split_param,
                "checkpoint": ckpt,
                "best_monitor": best_monitor,
                "best_mode": best_mode,
                "eval_mode": eval_mode,
                "n_folds": len(items),
                "folds": ",".join(str(x["fold"]) for x in items),
                "mean_dice": statistics.mean(dices),
                "std_dice": statistics.stdev(dices) if len(dices) > 1 else 0.0,
                "median_dice": statistics.median(dices),
                "best_dice": max(dices),
                "worst_dice": min(dices),
                "mean_nsd": statistics.mean(nsds) if nsds else None,
                "case_median_dice": statistics.median(case_dices) if case_dices else None,
                "case_iqr_low": percentile_or_none(case_dices, 25),
                "case_iqr_high": percentile_or_none(case_dices, 75),
                "zero_dice_cases": sum(1 for x in case_dices if x <= 0.0),
                "dice_ge_0_1_cases": sum(1 for x in case_dices if x >= 0.1),
                "dice_ge_0_2_cases": sum(1 for x in case_dices if x >= 0.2),
                "mean_unflagged": statistics.mean(unflagged) if unflagged else None,
                "mean_flagged": statistics.mean(flagged) if flagged else None,
                "n_flagged": sum(int(x["n_flagged"]) for x in items),
                "n_test": sum(int(x["n_test"]) for x in items),
                "key": key,
            }
        )
    return sorted(
        summary_rows,
        key=lambda x: (
            cohort_priority(str(x["cohort"])),
            -float(x["mean_dice"]),
            str(x["group"]),
            str(x["model"]),
        ),
    )


def append_header(ws, values: list[str], fill=HEADER_FILL) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.border = THIN
        cell.alignment = Alignment(horizontal="center", vertical="center")


def append_data(ws, values: list[Any]) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.border = THIN
        if isinstance(cell.value, float):
            cell.number_format = "0.000000"


def autosize(ws) -> None:
    widths: dict[str, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), len(str(cell.value)))
    for col, width in widths.items():
        ws.column_dimensions[col].width = min(max(width + 2, 12), 52)


def display_model(model_name: str) -> str:
    return MODEL_DISPLAY.get(model_name, model_name)


def model_architecture(model_name: str) -> str:
    info = MODEL_INFO.get(model_name)
    return info["architecture"] if info else ""


def comparison_id(row: dict[str, Any]) -> tuple[Any, ...]:
    return (
        row.get("key"),
        row.get("cohort"),
        row.get("group"),
        row.get("model"),
        row.get("setting"),
        row.get("split_method"),
        row.get("split_param"),
        row.get("checkpoint"),
        row.get("best_monitor"),
        row.get("best_mode"),
        row.get("eval_mode"),
    )


def find_comparison_row(summary_rows: list[dict[str, Any]], spec: dict[str, Any]) -> dict[str, Any] | None:
    for key in spec["keys"]:
        candidates = []
        for row in summary_rows:
            if row["key"] != key:
                continue
            if row["model"] != spec["model"]:
                continue
            if spec.get("eval_mode") is not None and row.get("eval_mode") != spec["eval_mode"]:
                continue
            if row.get("checkpoint") != "best":
                continue
            candidates.append(row)
        if candidates:
            return sorted(
                candidates,
                key=lambda x: (
                    -int(x["n_folds"]),
                    -float(x["mean_dice"]),
                    str(x.get("eval_mode")),
                ),
            )[0]
    return None


def selected_comparison_rows(summary_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    selected = []
    for order, spec in enumerate(COMPARISON_SPECS):
        row = find_comparison_row(summary_rows, spec)
        if row is None:
            row = {
                "group": model_group(spec["model"]),
                "cohort": spec["cohort"],
                "model": spec["model"],
                "setting": "",
                "split_method": "",
                "split_param": "",
                "checkpoint": "",
                "best_monitor": "",
                "best_mode": "",
                "eval_mode": spec.get("eval_mode"),
                "n_folds": 0,
                "folds": "",
                "mean_dice": None,
                "std_dice": None,
                "median_dice": None,
                "best_dice": None,
                "worst_dice": None,
                "mean_nsd": None,
                "case_median_dice": None,
                "case_iqr_low": None,
                "case_iqr_high": None,
                "zero_dice_cases": None,
                "dice_ge_0_1_cases": None,
                "dice_ge_0_2_cases": None,
                "mean_unflagged": None,
                "mean_flagged": None,
                "n_flagged": 0,
                "n_test": 0,
                "key": spec["keys"][0],
            }
        row = dict(row)
        row["display_model"] = display_model(str(row["model"]))
        row["architecture"] = model_architecture(str(row["model"]))
        row["comparison_cohort"] = spec["cohort"]
        row["expected_folds"] = spec["expected_folds"]
        row["comparison_order"] = order
        row["notes"] = comparison_note(row, spec)
        selected.append(row)
    return selected


def comparison_note(row: dict[str, Any], spec: dict[str, Any]) -> str:
    notes = [spec["note"]]
    if row["n_folds"] == 0:
        notes.append("No result JSON found yet.")
    elif int(row["n_folds"]) < int(spec["expected_folds"]):
        notes.append(f"Incomplete: {row['n_folds']}/{spec['expected_folds']} folds available.")
    if (
        "task10_qc76_baselines_20260517" in str(row["key"])
        and str(row["model"]) in {"unet_b", "unet_xl"}
    ):
        notes.append("Temporary older QC76 baseline; full 100-batch rerun is still pending.")
    if "task10_qc76_baselines_fulltrain_20260518" in str(row["key"]):
        notes.append("Full 100-batch QC76 baseline rerun.")
    if "task10_qc76_baselines_fulltrain_nomirror_argmax_20260519" in str(row["key"]):
        notes.append("Official-style eval-only rerun from the same fulltrain best checkpoints.")
    return " ".join(notes)


def write_title(ws, title: str, note: str | None = None) -> None:
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Updated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    if note:
        ws.append(["Note", note])
    ws.append([])


def write_report_title(ws, title: str, note: str | None = None) -> None:
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    if note:
        ws.append([note])
    else:
        ws.append([])


def write_metrics_sheet(wb, selected_rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("FCD_Seg_Metrics", 1)
    write_report_title(
        ws,
        "FCD segmentation comparison",
        "Held-out test-fold inference results. Original full-85 and QC76 flagged-excluded results are shown together for comparison.",
    )
    append_header(
        ws,
        [
            "Cohort",
            "Group",
            "Model",
            "Architecture",
            "Setting",
            "Split Method",
            "Checkpoint",
            "Eval/Postprocess",
            "N Folds",
            "Expected Folds",
            "Folds",
            "Total Test Predictions",
            "Mean Test Dice",
            "Std Test Dice",
            "Median Fold Dice",
            "Best Fold Dice",
            "Worst Fold Dice",
            "Mean Test NSD",
            "Median Case Dice",
            "Case Dice Q1",
            "Case Dice Q3",
            "Zero-Dice Cases",
            "Cases Dice >= 0.1",
            "Cases Dice >= 0.2",
            "Notes",
            "Result Key",
        ],
    )
    for row in selected_rows:
        append_data(
            ws,
            [
                row["comparison_cohort"],
                row["group"],
                row["display_model"],
                row["architecture"],
                row["setting"],
                row["split_method"],
                row["checkpoint"],
                row["eval_mode"],
                row["n_folds"],
                row["expected_folds"],
                row["folds"],
                row["n_test"],
                row["mean_dice"],
                row["std_dice"],
                row["median_dice"],
                row["best_dice"],
                row["worst_dice"],
                row["mean_nsd"],
                row["case_median_dice"],
                row["case_iqr_low"],
                row["case_iqr_high"],
                row["zero_dice_cases"],
                row["dice_ge_0_1_cases"],
                row["dice_ge_0_2_cases"],
                row["notes"],
                row["key"],
            ],
        )
    ws.freeze_panes = "A4"
    autosize(ws)


def selected_fold_rows(
    selected_rows: list[dict[str, Any]], fold_rows: list[dict[str, Any]]
) -> list[tuple[dict[str, Any], dict[str, Any]]]:
    selected_ids = {comparison_id(row): row for row in selected_rows if row.get("n_folds", 0)}
    rows = []
    for row in fold_rows:
        selected = selected_ids.get(comparison_id(row))
        if selected:
            rows.append((selected, row))
    return sorted(
        rows,
        key=lambda x: (
            int(x[0]["comparison_order"]),
            int(x[1]["fold"] or -1),
        ),
    )


def write_predictions_sheet(wb, selected_rows: list[dict[str, Any]], fold_rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("FCD_Seg_Predictions", 2)
    write_report_title(
        ws,
        "FCD segmentation held-out case predictions",
        "Each row is one held-out test case from one fold. Dice/NSD summarize the predicted mask against the ROI.",
    )
    append_header(
        ws,
        [
            "Cohort",
            "Subject",
            "Group",
            "Model",
            "Fold",
            "QC Flagged",
            "GT Voxels",
            "Checkpoint",
            "Eval/Postprocess",
            "Dice",
            "NSD",
            "Pred Voxels",
            "Result JSON",
        ],
    )
    for selected, row in selected_fold_rows(selected_rows, fold_rows):
        for pred in row.get("predictions", []):
            append_data(
                ws,
                [
                    selected["comparison_cohort"],
                    pred["id"],
                    row["group"],
                    display_model(str(row["model"])),
                    row["fold"],
                    pred["flagged_qc"],
                    pred["gt_voxels"],
                    row["checkpoint"],
                    row["eval_mode"],
                    pred["dice"],
                    pred["nsd"],
                    pred["pred_voxels"],
                    row["path"],
                ],
            )
    ws.freeze_panes = "A4"
    autosize(ws)


def paired_test(values_a: list[float], values_b: list[float]) -> dict[str, Any]:
    diffs = [a - b for a, b in zip(values_a, values_b)]
    out: dict[str, Any] = {
        "mean_a": statistics.mean(values_a),
        "mean_b": statistics.mean(values_b),
        "median_a": statistics.median(values_a),
        "median_b": statistics.median(values_b),
        "mean_diff": statistics.mean(diffs),
        "t_stat": None,
        "t_p": None,
        "wilcoxon_p": None,
    }
    try:
        from scipy import stats

        t_result = stats.ttest_rel(values_a, values_b)
        out["t_stat"] = None if math.isnan(float(t_result.statistic)) else float(t_result.statistic)
        out["t_p"] = None if math.isnan(float(t_result.pvalue)) else float(t_result.pvalue)
        if any(abs(x) > 0 for x in diffs):
            w_result = stats.wilcoxon(values_a, values_b, zero_method="zsplit")
            out["wilcoxon_p"] = None if math.isnan(float(w_result.pvalue)) else float(w_result.pvalue)
    except Exception:
        pass
    return out


def prediction_map(selected: dict[str, Any], fold_rows: list[dict[str, Any]]) -> dict[str, float]:
    selected_id = comparison_id(selected)
    out = {}
    for row in fold_rows:
        if comparison_id(row) != selected_id:
            continue
        for pred in row.get("predictions", []):
            out[str(pred["id"])] = float(pred["dice"])
    return out


def statistics_rows(
    selected_rows: list[dict[str, Any]], fold_rows: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    complete_rows = [r for r in selected_rows if r.get("n_folds", 0)]
    by_cohort: dict[str, list[dict[str, Any]]] = {}
    for row in complete_rows:
        by_cohort.setdefault(str(row["comparison_cohort"]), []).append(row)

    rows = []
    for cohort, items in by_cohort.items():
        items_by_model = {str(x["model"]): x for x in items}
        pairs: list[tuple[dict[str, Any], dict[str, Any]]] = []
        ours = items_by_model.get("cleandift_s23")
        if ours:
            pairs.extend((ours, x) for x in items if x is not ours)
        if "unet_b" in items_by_model and "unet_xl" in items_by_model:
            pairs.append((items_by_model["unet_b"], items_by_model["unet_xl"]))

        for left, right in pairs:
            left_map = prediction_map(left, fold_rows)
            right_map = prediction_map(right, fold_rows)
            subjects = sorted(set(left_map) & set(right_map))
            if not subjects:
                continue
            values_a = [left_map[s] for s in subjects]
            values_b = [right_map[s] for s in subjects]
            stats_out = paired_test(values_a, values_b)
            rows.append(
                {
                    "cohort": cohort,
                    "model_a": display_model(str(left["model"])),
                    "model_b": display_model(str(right["model"])),
                    "n": len(subjects),
                    **stats_out,
                }
            )
    threshold = 0.05 / len(rows) if rows else None
    for row in rows:
        row["bonferroni_threshold"] = threshold
        p = row.get("wilcoxon_p") if row.get("wilcoxon_p") is not None else row.get("t_p")
        row["significant_bonferroni"] = bool(p is not None and threshold is not None and p < threshold)
    return rows


def write_statistics_sheet(wb, selected_rows: list[dict[str, Any]], fold_rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("FCD_Seg_Statistics", 3)
    write_report_title(
        ws,
        "FCD segmentation paired statistics",
        "Paired tests use per-case held-out Dice values for models evaluated on the same subjects.",
    )
    append_header(
        ws,
        [
            "Cohort",
            "Model A",
            "Model B",
            "N Paired Cases",
            "Model A Mean Dice",
            "Model B Mean Dice",
            "Mean Dice Difference A-B",
            "Model A Median Dice",
            "Model B Median Dice",
            "Paired t-stat",
            "Paired t-test p",
            "Wilcoxon signed-rank p",
            "Bonferroni p threshold",
            "Significant After Bonferroni",
        ],
    )
    for row in statistics_rows(selected_rows, fold_rows):
        append_data(
            ws,
            [
                row["cohort"],
                row["model_a"],
                row["model_b"],
                row["n"],
                row["mean_a"],
                row["mean_b"],
                row["mean_diff"],
                row["median_a"],
                row["median_b"],
                row["t_stat"],
                row["t_p"],
                row["wilcoxon_p"],
                row["bonferroni_threshold"],
                row["significant_bonferroni"],
            ],
        )
    ws.freeze_panes = "A4"
    autosize(ws)


def write_model_info_sheet(wb) -> None:
    ws = wb.create_sheet("FCD_Seg_ModelInfo", 4)
    write_report_title(ws, "FCD segmentation model notes", COMMON_SEGMENTATION_SETUP)
    append_header(
        ws,
        [
            "Model",
            "Group",
            "Architecture Summary",
            "Segmentation Input",
            "Patch/Volume Setup",
            "Decoder/Head",
        ],
    )
    for model_name in ["cleandift_s23", "unet_b", "unet_xl", "mmunetvae"]:
        info = MODEL_INFO[model_name]
        append_data(
            ws,
            [
                display_model(model_name),
                info["group"],
                info["architecture"],
                info["input"],
                info["resolution"],
                info["decoder"],
            ],
        )
    ws.append([])
    ws.append(["QC76 excluded cases", ", ".join(sorted(FLAGGED_QC_CASES))])
    ws["A" + str(ws.max_row)].font = Font(bold=True)
    autosize(ws)


def write_unetxl_audit_sheet(wb, summary_rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("FCD_Seg_UNetXL_Audit", 5)
    write_report_title(
        ws,
        "FCD segmentation UNet-XL result audit",
        (
            "Historical UNet-XL rows are listed to show which results are complete and "
            "cohort-matched. Incomplete or cohort-mismatched rows should not replace the "
            "current QC76 comparison row."
        ),
    )
    append_header(
        ws,
        [
            "Use For QC76 Main Comparison?",
            "Reason",
            "Cohort",
            "Setting",
            "Split Method",
            "Checkpoint",
            "Eval/Postprocess",
            "N Folds",
            "Folds",
            "Total Test Predictions",
            "Mean Test Dice",
            "Median Case Dice",
            "Zero-Dice Cases",
            "Cases Dice >= 0.1",
            "Result Key",
        ],
    )
    rows = [r for r in summary_rows if r.get("model") == "unet_xl"]
    rows = sorted(
        rows,
        key=lambda x: (
            "qc76_unetxl_reseed_20260519" not in str(x["key"]),
            int(x["n_folds"]) != 5,
            float(x["mean_dice"]) if x["mean_dice"] is not None else 999,
            str(x["key"]),
        ),
    )
    for row in rows:
        key = str(row["key"])
        cohort = str(row["cohort"])
        n_folds = int(row["n_folds"])
        use = "No"
        reason = ""
        if key == (
            "task10_qc76_unetxl_reseed_20260519/"
            "unet_xl_qc76_unetxl_b4_all_dicece_best_100b_seed20260519_v1"
        ):
            use = "Yes"
            reason = "Complete 5-fold QC76 run paired with the current standardized UNet-B QC76 run."
        elif (
            "task10_qc76_baselines_20260517/"
            "unet_xl_qc76_unetxl_b4_all_dicece_best_nw2_v2" in key
        ):
            reason = (
                "QC76 split, but not reportable: short exploratory run "
                "(40 epochs, 25 train batches/epoch), only folds 0/2/4 completed, "
                "and fold-level evaluation mixed argmax and tuned settings."
            )
        elif cohort.startswith("FCD QC76") and n_folds < 5:
            reason = "QC76 cohort, but incomplete/exploratory fold coverage."
        elif cohort.startswith("FCD full"):
            reason = "Complete result, but full-85 cohort is not the QC76 flagged-excluded cohort."
        elif "baseline_default" in key:
            reason = "Complete sensitivity run, but separate baseline-default protocol."
        else:
            reason = "Historical or exploratory setting, not the selected QC76 main comparison."
        append_data(
            ws,
            [
                use,
                reason,
                row["cohort"],
                row["setting"],
                row["split_method"],
                row["checkpoint"],
                row["eval_mode"],
                row["n_folds"],
                row["folds"],
                row["n_test"],
                row["mean_dice"],
                row["case_median_dice"],
                row["zero_dice_cases"],
                row["dice_ge_0_1_cases"],
                row["key"],
            ],
        )
    ws.freeze_panes = "A4"
    autosize(ws)


def write_sheet(workbook: Path, summary_rows: list[dict[str, Any]], fold_rows: list[dict[str, Any]]) -> None:
    wb = load_workbook(workbook)
    for sheet_name in [MASTER_SHEET, *REPLACED_SHEETS]:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
    ws = wb.create_sheet(MASTER_SHEET, 0)

    ws.append(["FCD segmentation master result tracker"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(
        [
            "Updated",
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            "Publication-facing rows should use best.ckpt unless explicitly marked otherwise. QC76 excludes the 9 overlay-flagged cases.",
        ]
    )
    ws.append(["Flagged QC cases", ", ".join(sorted(FLAGGED_QC_CASES))])
    ws.append([])

    summary_header = [
        "Group",
        "Cohort",
        "Model",
        "Setting",
        "Split Method",
        "Split Param",
        "Checkpoint",
        "Best Monitor",
        "Best Mode",
        "Eval Mode",
        "N Folds",
        "Folds",
        "Mean Test Dice",
        "Std Test Dice Across Folds",
        "Median Fold Test Dice",
        "Best Fold Dice",
        "Worst Fold Dice",
        "Mean Test NSD",
        "Median Case Dice",
        "Case Dice Q1",
        "Case Dice Q3",
        "Zero-Dice Cases",
        "Cases Dice >= 0.1",
        "Cases Dice >= 0.2",
        "Mean Dice Unflagged Cases",
        "Mean Dice Flagged Cases",
        "Flagged Test Cases Total",
        "Total Test Predictions",
        "Result Key",
    ]
    append_header(ws, summary_header)
    for row in summary_rows:
        append_data(
            ws,
            [
                row["group"],
                row["cohort"],
                row["model"],
                row["setting"],
                row["split_method"],
                row["split_param"],
                row["checkpoint"],
                row["best_monitor"],
                row["best_mode"],
                row["eval_mode"],
                row["n_folds"],
                row["folds"],
                row["mean_dice"],
                row["std_dice"],
                row["median_dice"],
                row["best_dice"],
                row["worst_dice"],
                row["mean_nsd"],
                row["case_median_dice"],
                row["case_iqr_low"],
                row["case_iqr_high"],
                row["zero_dice_cases"],
                row["dice_ge_0_1_cases"],
                row["dice_ge_0_2_cases"],
                row["mean_unflagged"],
                row["mean_flagged"],
                row["n_flagged"],
                row["n_test"],
                row["key"],
            ],
        )

    ws.append([])
    ws.append(["Per-fold details"])
    ws[ws.max_row][0].font = Font(bold=True, size=12)
    append_header(
        ws,
        [
            "Group",
            "Cohort",
            "Model",
            "Setting",
            "Split Method",
            "Split Param",
            "Checkpoint",
            "Best Monitor",
            "Best Mode",
            "Eval Mode",
            "Fold",
            "Test Dice",
            "Test NSD",
            "Median Case Dice",
            "Case Dice Q1",
            "Case Dice Q3",
            "Zero-Dice Cases",
            "Cases Dice >= 0.1",
            "Cases Dice >= 0.2",
            "Mean Dice Unflagged Cases",
            "Mean Dice Flagged Cases",
            "N Flagged Test Cases",
            "N Test Cases",
            "Result JSON",
        ],
        fill=SUBHEADER_FILL,
    )
    for row in sorted(
        fold_rows,
        key=lambda x: (
            cohort_priority(str(x["cohort"])),
            str(x["group"]),
            str(x["model"]),
            str(x["key"]),
            int(x["fold"] or -1),
        ),
    ):
        append_data(
            ws,
            [
                row["group"],
                row["cohort"],
                row["model"],
                row["setting"],
                row["split_method"],
                row["split_param"],
                row["checkpoint"],
                row["best_monitor"],
                row["best_mode"],
                row["eval_mode"],
                row["fold"],
                row["test_dice"],
                row["test_nsd"],
                row["case_median_dice"],
                row["case_iqr_low"],
                row["case_iqr_high"],
                row["zero_dice_cases"],
                row["dice_ge_0_1_cases"],
                row["dice_ge_0_2_cases"],
                row["unflagged_dice"],
                row["flagged_dice"],
                row["n_flagged"],
                row["n_test"],
                row["path"],
            ],
        )

    ws.freeze_panes = "A6"
    autosize(ws)
    selected_rows = selected_comparison_rows(summary_rows)
    write_metrics_sheet(wb, selected_rows)
    write_predictions_sheet(wb, selected_rows, fold_rows)
    write_statistics_sheet(wb, selected_rows, fold_rows)
    write_model_info_sheet(wb)
    write_unetxl_audit_sheet(wb, summary_rows)
    wb.save(workbook)


def main() -> None:
    args = parse_args()
    rows = collect_rows(args)
    if not rows:
        raise SystemExit("No segmentation result JSONs found.")
    write_sheet(args.workbook, summarize(rows), rows)
    print(f"Updated {args.workbook} [{MASTER_SHEET}] with {len(rows)} fold rows.")


if __name__ == "__main__":
    main()
