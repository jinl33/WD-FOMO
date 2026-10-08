#!/usr/bin/env python3

from __future__ import annotations

import argparse
import json
import re
import shlex
import shutil
import statistics
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

DEFAULT_WORKBOOK = Path("/nfs/turbo/umms-wilms1/FOMO/inference_outputs/FOMO300K_results_v4.xlsx")
DEFAULT_LOG_DIR = Path("/nfs/turbo/umms-wilms1/FOMO/experiments/jinhlee/logs")

SHEETS = ["ATLAS_Seg_Metrics", "ATLAS_Seg_Predictions", "ATLAS_Seg_Statistics", "ATLAS_Seg_Timing"]

MODEL_GROUP = {
    "Ours-23M": "Ours",
    "UNet-B": "Baseline",
    "UNet-XL": "Baseline",
    "MMUNetVAE": "Baseline",
}

SPECS = [
    {
        "cohort": "955-case single split (8:1:1 stratified, test n=95)",
        "input": "T1 only",
        "model": "Ours-23M",
        "status": "Complete",
        "expected_folds": 1,
        "paths": [
            "/nfs/turbo/umms-wilms1/FOMO/inference_outputs/task12_single_run_all955_20260528/"
            "cleandift_s23_atlas955_single80_10_10_waveletavg_refine_fold0_test.json",
        ],
    },
    {
        "cohort": "955-case single split (8:1:1 stratified, test n=95)",
        "input": "T1 only",
        "model": "UNet-B",
        "status": "Complete",
        "expected_folds": 1,
        "paths": [
            "/nfs/turbo/umms-wilms1/FOMO/inference_outputs/task12_single_run_all955_20260531_recovered_eval/"
            "unet_b_atlas955_single80_10_10_unetb_repohparamslocal_fold0_best_test.json",
        ],
    },
    {
        "cohort": "955-case single split (8:1:1 stratified, test n=95)",
        "input": "T1 only",
        "model": "UNet-XL",
        "status": "Complete",
        "expected_folds": 1,
        "paths": [
            "/nfs/turbo/umms-wilms1/FOMO/inference_outputs/task12_single_run_all955_20260531_recovered_eval/"
            "unet_xl_atlas955_single80_10_10_unetxl_repohparamslocal_fold0_best_test.json",
        ],
    },
    {
        "cohort": "955-case single split (8:1:1 stratified, test n=95)",
        "input": "T1 only",
        "model": "MMUNetVAE",
        "status": "Complete",
        "expected_folds": 1,
        "paths": [
            "/nfs/turbo/umms-wilms1/FOMO/inference_outputs/task12_single_run_all955_20260528/"
            "mmunetvae_atlas955_single80_10_10_mmunetvae_nomirror_last_fold0_test.json",
        ],
    },
    {
        "cohort": "955-case paper-style 10-center holdout (test n=405)",
        "input": "T1 only",
        "model": "Ours-23M",
        "status": "Complete",
        "expected_folds": 1,
        "paths": [
            "/nfs/turbo/umms-wilms1/FOMO/inference_outputs/task12_paper10center_all955_20260531_recovered_eval/"
            "cleandift_s23_atlas955_paper10center_seed42_waveletavg_refine_fold0_best_test.json",
        ],
    },
    {
        "cohort": "955-case paper-style 10-center holdout (test n=405)",
        "input": "T1 only",
        "model": "MMUNetVAE",
        "status": "Complete",
        "expected_folds": 1,
        "paths": [
            "/nfs/turbo/umms-wilms1/FOMO/inference_outputs/task12_paper10center_all955_20260531_recovered_eval/"
            "mmunetvae_atlas955_paper10center_seed42_mmunetvae_fold0_best_test.json",
        ],
    },
    {
        "cohort": "955-case paper-style 10-center holdout (test n=405)",
        "input": "T1 only",
        "model": "UNet-B",
        "status": "Complete",
        "expected_folds": 1,
        "paths": [
            "/nfs/turbo/umms-wilms1/FOMO/inference_outputs/task12_paper10center_all955_20260531_recovered_eval/"
            "unet_b_atlas955_paper10center_seed42_unetb_fold0_best_test.json",
        ],
    },
    {
        "cohort": "955-case paper-style 10-center holdout (test n=405)",
        "input": "T1 only",
        "model": "UNet-XL",
        "status": "Complete",
        "expected_folds": 1,
        "paths": [
            "/nfs/turbo/umms-wilms1/FOMO/inference_outputs/task12_paper10center_all955_20260531_recovered_eval/"
            "unet_xl_atlas955_paper10center_seed42_unetxl_fold0_best_test.json",
        ],
    },
]

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser()
    p.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    p.add_argument("--log-dir", type=Path, default=DEFAULT_LOG_DIR)
    return p.parse_args()

def append_title(ws, title: str, subtitle: str, width: int) -> None:
    ws.append([title])
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.append([subtitle])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws.append([])

def append_header(ws, header: list[str]) -> None:
    ws.append(header)
    fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

def autosize(ws) -> None:
    for column_cells in ws.columns:
        first_cell = next((cell for cell in column_cells if hasattr(cell, "column_letter")), None)
        if first_cell is None:
            continue
        letter = first_cell.column_letter
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 80))
        ws.column_dimensions[letter].width = max(10, min(max_len + 2, 60))

def optional_float(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except Exception:
        return None

def fmt_hms(seconds: float | None) -> str:
    if seconds is None:
        return ""
    total = int(round(float(seconds)))
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"

def _grep_logs_for_literal(log_dir: Path, literal: str, max_hits: int = 20) -> list[Path]:
    if not log_dir.exists() or not literal:
        return []
    cmd = (
        f"grep -lF -- {shlex.quote(literal)} {shlex.quote(str(log_dir))}/*.log 2>/dev/null "
        f"| head -n {int(max_hits)}"
    )
    try:
        proc = subprocess.run(["bash", "-lc", cmd], check=False, capture_output=True, text=True)
    except Exception:
        return []
    out = (proc.stdout or "").strip()
    if not out:
        return []
    return [Path(line.strip()) for line in out.splitlines() if line.strip()]

def _parse_training_timing(log_text: str) -> dict[str, Any]:
    timing: dict[str, Any] = {
        "finetune_seconds": None,
        "epochs_ran": None,
        "seconds_per_epoch": None,
        "best_epoch": None,
        "best_val_dice": None,
        "seconds_to_best": None,
    }

    done_match = re.search(r"Fine-tune done:\s*([0-9]+)s", log_text)
    if done_match:
        timing["finetune_seconds"] = float(done_match.group(1))

    epoch_re = re.compile(
        r"Epoch\s+(\d+):\s+100%.*?\[(\d+):(\d+)<.*?val/dice=([0-9]*\.?[0-9]+)",
        re.DOTALL,
    )

    per_epoch_seconds: dict[int, float] = {}
    per_epoch_val: dict[int, float] = {}
    for m in epoch_re.finditer(log_text):
        epoch = int(m.group(1))
        dur_sec = int(m.group(2)) * 60 + int(m.group(3))
        val_dice = float(m.group(4))
        per_epoch_seconds.setdefault(epoch, float(dur_sec))
        prev = per_epoch_val.get(epoch)
        per_epoch_val[epoch] = val_dice if prev is None else max(prev, val_dice)

    if per_epoch_val:
        max_epoch = max(per_epoch_val)
        timing["epochs_ran"] = int(max_epoch + 1)
        best_epoch = max(per_epoch_val, key=lambda e: per_epoch_val[e])
        timing["best_epoch"] = int(best_epoch)
        timing["best_val_dice"] = float(per_epoch_val[best_epoch])

        if timing["finetune_seconds"] is not None and timing["epochs_ran"]:
            timing["seconds_per_epoch"] = float(timing["finetune_seconds"]) / float(timing["epochs_ran"])

        if per_epoch_seconds and best_epoch is not None:
            got_all = all(e in per_epoch_seconds for e in range(best_epoch + 1))
            if got_all:
                timing["seconds_to_best"] = float(sum(per_epoch_seconds[e] for e in range(best_epoch + 1)))

        if timing["seconds_to_best"] is None and timing["seconds_per_epoch"] is not None and best_epoch is not None:
            timing["seconds_to_best"] = float(timing["seconds_per_epoch"]) * float(best_epoch + 1)

    return timing

def _select_training_log(
    log_dir: Path,
    output_json_path: str,
    checkpoint_path: str,
) -> Path | None:
    candidates: list[Path] = []
    candidates.extend(_grep_logs_for_literal(log_dir, f"Output JSON: {output_json_path}"))
    candidates.extend(_grep_logs_for_literal(log_dir, f"Checkpoint: {checkpoint_path}"))
    dedup: list[Path] = []
    seen = set()
    for p in candidates:
        if p in seen:
            continue
        seen.add(p)
        dedup.append(p)
    if not dedup:
        return None

    for p in dedup:
        try:
            txt = p.read_text(errors="ignore")
        except Exception:
            continue
        if "Fine-tune done:" in txt:
            return p
    return dedup[0]

def load_predictions(spec: dict[str, Any]) -> tuple[list[dict[str, Any]], list[int]]:
    rows: list[dict[str, Any]] = []
    folds: list[int] = []
    for raw_path in spec["paths"]:
        path = Path(raw_path)
        payload = json.loads(path.read_text())
        fold = int(payload["fold_idx"])
        folds.append(fold)
        preds = payload.get("predictions", payload.get("subjects", []))
        for pred in preds:
            rows.append(
                {
                    "cohort": spec["cohort"],
                    "group": MODEL_GROUP[spec["model"]],
                    "model": spec["model"],
                    "input": spec["input"],
                    "status": spec["status"],
                    "fold": fold,
                    "subject": pred["id"],
                    "dice": float(pred["dice"]),
                    "nsd": optional_float(pred.get("nsd")),
                    "gt_voxels": pred.get("gt_voxels"),
                    "pred_voxels": pred.get("pred_voxels"),
                    "flagged_qc": "",
                    "result_json": str(path),
                }
            )
    return rows, folds

def collect_rows(log_dir: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    metrics: list[dict[str, Any]] = []
    predictions: list[dict[str, Any]] = []
    timing_rows: list[dict[str, Any]] = []

    for spec in SPECS:
        pred_rows, folds = load_predictions(spec)
        predictions.extend(pred_rows)
        per_fold_timing: list[dict[str, Any]] = []
        for raw_path in spec["paths"]:
            payload = json.loads(Path(raw_path).read_text())
            fold = int(payload["fold_idx"])
            checkpoint_path = str(payload.get("checkpoint", ""))
            train_log = _select_training_log(log_dir, raw_path, checkpoint_path)
            timing = {
                "finetune_seconds": None,
                "epochs_ran": None,
                "seconds_per_epoch": None,
                "best_epoch": None,
                "best_val_dice": None,
                "seconds_to_best": None,
            }
            if train_log is not None:
                try:
                    timing = _parse_training_timing(train_log.read_text(errors="ignore"))
                except Exception:
                    pass
            detail = {
                "cohort": spec["cohort"],
                "group": MODEL_GROUP[spec["model"]],
                "model": spec["model"],
                "input": spec["input"],
                "status": spec["status"],
                "fold": fold,
                "training_log": str(train_log) if train_log else "",
                "checkpoint": checkpoint_path,
            } | timing
            timing_rows.append(detail)
            per_fold_timing.append(detail)

        complete = [row for row in pred_rows if row["dice"] is not None]
        dice = [row["dice"] for row in complete]
        nsd = [row["nsd"] for row in complete if row["nsd"] is not None]
        subjects = sorted({row["subject"] for row in complete})

        ft_vals = [float(r["finetune_seconds"]) for r in per_fold_timing if r["finetune_seconds"] is not None]
        ep_vals = [float(r["seconds_per_epoch"]) for r in per_fold_timing if r["seconds_per_epoch"] is not None]
        best_ep_vals = [float(r["best_epoch"]) for r in per_fold_timing if r["best_epoch"] is not None]
        best_time_vals = [float(r["seconds_to_best"]) for r in per_fold_timing if r["seconds_to_best"] is not None]

        metrics.append(
            {
                "cohort": spec["cohort"],
                "group": MODEL_GROUP[spec["model"]],
                "model": spec["model"],
                "input": spec["input"],
                "status": spec["status"],
                "n_subjects": len(subjects),
                "n_folds": len(set(folds)),
                "expected_folds": int(spec["expected_folds"]),
                "mean_dice": statistics.mean(dice) if dice else None,
                "std_dice": statistics.pstdev(dice) if len(dice) > 1 else 0.0 if dice else None,
                "median_dice": statistics.median(dice) if dice else None,
                "mean_nsd": statistics.mean(nsd) if nsd else None,
                "std_nsd": statistics.pstdev(nsd) if len(nsd) > 1 else 0.0 if nsd else None,
                "median_nsd": statistics.median(nsd) if nsd else None,
                "mean_finetune_seconds": statistics.mean(ft_vals) if ft_vals else None,
                "mean_seconds_per_epoch": statistics.mean(ep_vals) if ep_vals else None,
                "mean_best_epoch": statistics.mean(best_ep_vals) if best_ep_vals else None,
                "mean_seconds_to_best": statistics.mean(best_time_vals) if best_time_vals else None,
            }
        )
    predictions.sort(key=lambda r: (r["cohort"], r["input"], r["group"], r["model"], r["fold"], r["subject"]))
    metrics.sort(key=lambda r: (r["cohort"], r["input"], r["group"], r["model"]))
    timing_rows.sort(key=lambda r: (r["cohort"], r["input"], r["group"], r["model"], r["fold"]))
    return metrics, predictions, timing_rows

def wilcoxon_p(a: list[float], b: list[float]) -> float | None:
    try:
        from scipy.stats import wilcoxon
    except Exception:
        return None
    try:
        stat = wilcoxon(np.asarray(a), np.asarray(b), zero_method="zsplit")
    except ValueError:
        return None
    return float(stat.pvalue)

def paired_stats(predictions: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    by_group: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for row in predictions:
        by_group.setdefault((row["cohort"], row["input"]), []).append(row)
    for (cohort, input_name), items in sorted(by_group.items()):
        models = sorted({row["model"] for row in items})
        if "Ours-23M" not in models:
            continue
        for other in models:
            if other == "Ours-23M":
                continue
            for metric in ("dice", "nsd"):
                left = {row["subject"]: row[metric] for row in items if row["model"] == "Ours-23M" and row[metric] is not None}
                right = {row["subject"]: row[metric] for row in items if row["model"] == other and row[metric] is not None}
                subjects = sorted(set(left) & set(right))
                if not subjects:
                    continue
                a = [float(left[s]) for s in subjects]
                b = [float(right[s]) for s in subjects]
                diffs = [x - y for x, y in zip(a, b)]
                rows.append(
                    {
                        "cohort": cohort,
                        "input": input_name,
                        "metric": metric.upper(),
                        "comparison": f"Ours-23M - {other}",
                        "n": len(subjects),
                        "mean_ours": statistics.mean(a),
                        "mean_other": statistics.mean(b),
                        "mean_difference": statistics.mean(diffs),
                        "wilcoxon_p": wilcoxon_p(a, b),
                    }
                )
    return rows

def write_workbook(
    workbook: Path,
    metrics: list[dict[str, Any]],
    predictions: list[dict[str, Any]],
    timing_rows: list[dict[str, Any]],
) -> None:
    backup = workbook.with_suffix(f".backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    shutil.copy2(workbook, backup)
    wb = load_workbook(workbook)
    for sheet in SHEETS:
        if sheet in wb.sheetnames:
            del wb[sheet]

    ws = wb.create_sheet("ATLAS_Seg_Metrics")
    append_title(
        ws,
        "ATLAS segmentation metrics",
        "Held-out test performance on the current 955-case evaluations. Test-only reporting, aligned with the other downstream tasks.",
        14,
    )
    append_header(
        ws,
        [
            "Cohort",
            "Group",
            "Model",
            "Input",
            "Status",
            "N Subjects",
            "N Folds",
            "Expected Folds",
            "Mean Dice",
            "Std Dice",
            "Median Dice",
            "Mean NSD",
            "Std NSD",
            "Median NSD",
        ],
    )
    for row in metrics:
        ws.append(
            [
                row["cohort"],
                row["group"],
                row["model"],
                row["input"],
                row["status"],
                row["n_subjects"],
                row["n_folds"],
                row["expected_folds"],
                row["mean_dice"],
                row["std_dice"],
                row["median_dice"],
                row["mean_nsd"],
                row["std_nsd"],
                row["median_nsd"],
            ]
        )
    autosize(ws)

    ws = wb.create_sheet("ATLAS_Seg_Predictions")
    append_title(ws, "ATLAS segmentation predictions", "Held-out test case Dice/NSD.", 12)
    append_header(
        ws,
        [
            "Cohort",
            "Group",
            "Model",
            "Input",
            "Status",
            "Fold",
            "Subject",
            "Dice",
            "NSD",
            "GT Voxels",
            "Pred Voxels",
            "Flagged QC",
        ],
    )
    for row in predictions:
        ws.append(
            [
                row["cohort"],
                row["group"],
                row["model"],
                row["input"],
                row["status"],
                row["fold"],
                row["subject"],
                row["dice"],
                row["nsd"],
                row["gt_voxels"],
                row["pred_voxels"],
                row["flagged_qc"],
            ]
        )
    autosize(ws)

    stats = paired_stats(predictions)
    ws = wb.create_sheet("ATLAS_Seg_Statistics")
    append_title(ws, "ATLAS segmentation statistics", "Paired per-case comparisons on shared held-out test subjects.", 9)
    append_header(
        ws,
        [
            "Cohort",
            "Input",
            "Metric",
            "Comparison",
            "N",
            "Mean Ours",
            "Mean Other",
            "Mean Difference",
            "Wilcoxon p",
        ],
    )
    for row in stats:
        ws.append(
            [
                row["cohort"],
                row["input"],
                row["metric"],
                row["comparison"],
                row["n"],
                row["mean_ours"],
                row["mean_other"],
                row["mean_difference"],
                row["wilcoxon_p"],
            ]
        )
    autosize(ws)

    wb.save(workbook)
    print(f"Updated {workbook}")
    print(f"Backup {backup}")

def main() -> None:
    args = parse_args()
    metrics, predictions, timing_rows = collect_rows(args.log_dir)
    write_workbook(args.workbook, metrics, predictions, timing_rows)
    for row in metrics:
        print(
            row["cohort"],
            row["model"],
            row["status"],
            row["n_folds"],
            row["mean_dice"],
            row["mean_nsd"],
            fmt_hms(row.get("mean_finetune_seconds")),
            row.get("mean_best_epoch"),
        )

if __name__ == "__main__":
    main()
